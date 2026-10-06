"""
Comprehensive tests for Document Viewer, DOCX preview conversion, coordinate mapping,
and security access controls (Issues 2A - 2E).
"""
import io
import os
import fitz
import pytest
from unittest.mock import MagicMock, patch

from docx import Document
from tests.conftest import register_and_login
from app.services.document_viewer import (
    _find_ordered_spans_in_pdf,
    get_submission_document_pages,
)
from app.services.docx_converter import (
    DocxConversionError,
    get_converted_docx_pdf_path,
)
from app.services.extraction import extract_from_docx


def make_test_pdf_with_distinct_pages() -> bytes:
    """Builds a 3-page PDF with precise text locations."""
    doc = fitz.open()
    # Page 1
    p1 = doc.new_page(width=612, height=792)
    p1.insert_text((72, 100), "Title: Distributed Computing Protocols", fontsize=14)
    p1.insert_text((72, 150), "Primary analysis paragraph describing consensus mechanisms.", fontsize=11)

    # Page 2
    p2 = doc.new_page(width=612, height=792)
    p2.insert_text((72, 100), "Section 2: Replication Strategies and Fault Domains.", fontsize=12)
    p2.insert_text((72, 150), "Secondary analysis paragraph detailing state machine replication.", fontsize=11)

    # Page 3
    p3 = doc.new_page(width=612, height=792)
    p3.insert_text((72, 100), "Conclusion and references for academic review.", fontsize=11)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def make_test_docx_with_tables() -> bytes:
    """Builds a DOCX containing both paragraphs and a structured table."""
    doc = Document()
    doc.add_heading("Cloud Systems Evaluation", level=1)
    doc.add_paragraph("This paragraph introduces the cloud architectural benchmarks and network topologies.")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Observed Value"
    table.cell(1, 0).text = "Latency"
    table.cell(1, 1).text = "12 milliseconds"

    doc.add_paragraph("Concluding observations regarding system scalability and failover times.")

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


class TestOriginalBytesPreservation:
    """Tests that original uploaded PDF bytes are never corrupted or altered."""

    def test_original_bytes_intact(self, client):
        fac_headers = register_and_login(client, "faculty", "fac_bytes@example.com")["headers"]
        stu_headers = register_and_login(client, "student", "stu_bytes@example.com")["headers"]

        res = client.post(
            "/api/assignments",
            json={"title": "Bytes Preservation Assignment", "description": "Verify original bytes"},
            headers=fac_headers,
        )
        assignment_id = res.json()["id"]

        original_bytes = make_test_pdf_with_distinct_pages()
        upload_res = client.post(
            f"/api/assignments/{assignment_id}/submissions",
            files={"file": ("original_test.pdf", io.BytesIO(original_bytes), "application/pdf")},
            headers=stu_headers,
        )
        assert upload_res.status_code == 201
        submission_id = upload_res.json()["id"]

        # Fetch original file via /file endpoint
        file_res = client.get(f"/api/submissions/{submission_id}/file", headers=stu_headers)
        assert file_res.status_code == 200
        assert file_res.content == original_bytes, "Stored original PDF bytes do not match uploaded bytes!"


class TestNoMisleadingFallbackBox:
    """Tests that the legacy 8-92% fake fallback bounding box has been completely eliminated."""

    def test_phrase_not_found_returns_empty_list(self):
        pdf_bytes = make_test_pdf_with_distinct_pages()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")

        # Search for a phrase that definitely does NOT exist in the PDF
        non_existent_text = "Xylophone quantum entanglement zebra kaleidoscope non-existent phrase"
        spans_by_page = _find_ordered_spans_in_pdf(doc, non_existent_text)

        # Must return empty dict, NOT a fabricated whole-page box!
        assert spans_by_page == {}
        doc.close()


class TestAuthorizationAndNotFoundErrors:
    """Tests strict 403 authorization and 404 error cases for viewer endpoints."""

    def test_viewer_security_and_404s(self, client):
        fac1 = register_and_login(client, "faculty", "fac1_sec@example.com")["headers"]
        fac2 = register_and_login(client, "faculty", "fac2_sec@example.com")["headers"]
        stu1 = register_and_login(client, "student", "stu1_sec@example.com")["headers"]
        stu2 = register_and_login(client, "student", "stu2_sec@example.com")["headers"]

        res = client.post(
            "/api/assignments",
            json={"title": "Security Assignment", "description": "Auth testing"},
            headers=fac1,
        )
        assignment_id = res.json()["id"]

        pdf_bytes = make_test_pdf_with_distinct_pages()
        upload_res = client.post(
            f"/api/assignments/{assignment_id}/submissions",
            files={"file": ("sec.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
            headers=stu1,
        )
        submission_id = upload_res.json()["id"]

        # Non-existent submission UUID -> 404
        non_existent_uuid = "00000000-0000-0000-0000-000000000000"
        assert client.get(f"/api/submissions/{non_existent_uuid}/document-view", headers=stu1).status_code == 404
        assert client.get(f"/api/submissions/{non_existent_uuid}/pages/1/image", headers=stu1).status_code == 404
        assert client.get(f"/api/submissions/{non_existent_uuid}/annotated-file", headers=stu1).status_code == 404

        # Other student -> 403
        assert client.get(f"/api/submissions/{submission_id}/document-view", headers=stu2).status_code == 403
        assert client.get(f"/api/submissions/{submission_id}/pages/1/image", headers=stu2).status_code == 403
        assert client.get(f"/api/submissions/{submission_id}/annotated-file", headers=stu2).status_code == 403

        # Unrelated faculty (fac2 did not create assignment) -> 403
        assert client.get(f"/api/submissions/{submission_id}/document-view", headers=fac2).status_code == 403
        assert client.get(f"/api/submissions/{submission_id}/pages/1/image", headers=fac2).status_code == 403

        # Invalid page index (page 0 or page 99) -> 404
        assert client.get(f"/api/submissions/{submission_id}/pages/0/image", headers=stu1).status_code == 404
        assert client.get(f"/api/submissions/{submission_id}/pages/99/image", headers=stu1).status_code == 404


class TestDocxExtractionAndConversion:
    """Tests extraction of DOCX tables and LibreOffice conversion sandbox."""

    def test_docx_table_extraction(self, tmp_path):
        docx_bytes = make_test_docx_with_tables()
        temp_file = tmp_path / "table_test.docx"
        temp_file.write_bytes(docx_bytes)

        result = extract_from_docx(str(temp_file))
        assert result.error is None
        text = result.text

        # Must extract paragraph text
        assert "Cloud Systems Evaluation" in text
        assert "benchmarks and network topologies" in text

        # Must extract table content
        assert "Metric" in text
        assert "Latency" in text
        assert "12 milliseconds" in text

    @patch("shutil.which", return_value="/usr/bin/soffice")
    @patch("subprocess.run")
    def test_docx_conversion_sandbox_arguments(self, mock_subprocess, mock_which, tmp_path):
        docx_bytes = make_test_docx_with_tables()
        temp_docx = tmp_path / "sample.docx"
        temp_docx.write_bytes(docx_bytes)

        fake_pdf = b"%PDF-1.4 fake generated pdf bytes"

        def fake_run(cmd, **kwargs):
            # Verify command is a list without shell=True
            assert isinstance(cmd, list)
            assert "soffice" in cmd[0] or "libreoffice" in cmd[0]
            assert "--headless" in cmd
            assert "--convert-to" in cmd
            assert "pdf" in cmd
            assert kwargs.get("shell") is not True

            # Write fake PDF to output dir
            out_dir = cmd[cmd.index("--outdir") + 1]
            pdf_path = os.path.join(out_dir, f"{temp_docx.stem}.pdf")
            with open(pdf_path, "wb") as f:
                f.write(fake_pdf)
            mock_res = MagicMock()
            mock_res.returncode = 0
            return mock_res

        mock_subprocess.side_effect = fake_run

        pdf_path = get_converted_docx_pdf_path(str(temp_docx))
        assert os.path.exists(pdf_path)
        with open(pdf_path, "rb") as f:
            assert f.read() == fake_pdf


class TestCleanUnannotatedImageRetrieval:
    """Tests that annotated=false returns unannotated page image for frontend interactive overlays."""

    def test_unannotated_page_image_clean(self, client):
        fac_headers = register_and_login(client, "faculty", "fac_clean@example.com")["headers"]
        stu_headers = register_and_login(client, "student", "stu_clean@example.com")["headers"]

        res = client.post(
            "/api/assignments",
            json={"title": "Clean Image Assignment", "description": "Overlay testing"},
            headers=fac_headers,
        )
        assignment_id = res.json()["id"]

        pdf_bytes = make_test_pdf_with_distinct_pages()
        upload_res = client.post(
            f"/api/assignments/{assignment_id}/submissions",
            files={"file": ("clean.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
            headers=stu_headers,
        )
        submission_id = upload_res.json()["id"]

        clean_img_res = client.get(
            f"/api/submissions/{submission_id}/pages/1/image?annotated=false",
            headers=stu_headers,
        )
        assert clean_img_res.status_code == 200
        assert clean_img_res.headers["content-type"] == "image/png"
        assert len(clean_img_res.content) > 1000
